"""Figure replication: read source data, infer plot types, generate matplotlib figures."""

from __future__ import annotations

import csv
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import numpy as np

from .log import get_logger

logger = get_logger("figgen")

try:
    import openpyxl
except ImportError:
    openpyxl = None  # type: ignore[assignment]


@dataclass
class PanelSpec:
    sheet_name: str
    figure_number: str
    panel_label: str
    columns: list[str]
    data: list[dict]
    plot_type: str
    x_label: str = ""
    y_label: str = ""
    title: str = ""
    has_error_bars: bool = False


@dataclass
class FigureReplication:
    figure_number: str
    source_file: str
    panels: list[PanelSpec] = field(default_factory=list)
    output_paths: list[str] = field(default_factory=list)
    status: str = "pending"
    message: str = ""


def discover_source_data(package_dir: str | Path) -> dict[str, list[Path]]:
    """Map figure numbers to source data files. Prefers xlsx; groups CSVs by figure."""
    package_dir = Path(package_dir)
    xlsx_map: dict[str, Path] = {}
    csv_map: dict[str, list[Path]] = defaultdict(list)

    for p in sorted(package_dir.rglob("*")):
        if not p.is_file():
            continue
        name = p.name
        fig_num = _extract_fig_number(name)
        if not fig_num:
            continue

        if p.suffix.lower() in (".xlsx", ".xls"):
            xlsx_map[fig_num] = p
        elif p.suffix.lower() == ".csv":
            csv_map[fig_num].append(p)

    result: dict[str, list[Path]] = {}
    all_figs = set(xlsx_map.keys()) | set(csv_map.keys())
    for fig_num in all_figs:
        if fig_num in xlsx_map:
            result[fig_num] = [xlsx_map[fig_num]]
        elif fig_num in csv_map:
            result[fig_num] = sorted(csv_map[fig_num])

    return result


def _extract_fig_number(filename: str) -> str | None:
    """Extract figure number from a filename like SourceData_Fig1.xlsx or Supplementary_FigS5.csv."""
    name = filename.lower()
    m = re.search(r"sourcedata[_\-]?(ext(?:ended)?)?[_\-]?fig\.?(\d+)", name)
    if m:
        prefix = "E" if m.group(1) else ""
        return f"{prefix}{m.group(2)}"
    m = re.search(r"supplementary[_\-]?figs?\.?(\d+)", name)
    if m:
        return f"S{m.group(1)}"
    return None


def read_source_panels(files: list[Path], figure_number: str) -> list[PanelSpec]:
    """Read panels from xlsx (multi-sheet) or CSV files (one panel per file)."""
    if len(files) == 1 and files[0].suffix.lower() in (".xlsx", ".xls"):
        return _read_excel_panels(files[0], figure_number)
    return _read_csv_panels(files, figure_number)


def _read_excel_panels(filepath: Path, figure_number: str) -> list[PanelSpec]:
    """Read all sheets from an Excel source data file and classify each as a panel."""
    if openpyxl is None:
        raise ImportError("openpyxl required: pip install openpyxl")

    wb = openpyxl.load_workbook(str(filepath), read_only=True, data_only=True)
    panels = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            continue

        headers = [str(h) if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
        data = []
        for row in rows[1:]:
            if all(v is None for v in row):
                continue
            data.append(dict(zip(headers, row)))

        if not data:
            continue

        panel_label = _extract_panel_label(sheet_name)
        has_errors = any("err" in h.lower() or "lower" in h.lower() or "upper" in h.lower()
                        for h in headers)
        plot_type = _infer_plot_type(headers, data, sheet_name)

        panels.append(PanelSpec(
            sheet_name=sheet_name,
            figure_number=figure_number,
            panel_label=panel_label,
            columns=headers,
            data=data,
            plot_type=plot_type,
            has_error_bars=has_errors,
        ))

    wb.close()
    return panels


def _read_csv_panels(files: list[Path], figure_number: str) -> list[PanelSpec]:
    """Read CSV files, one per panel. Panel label extracted from filename."""
    panels = []
    for filepath in sorted(files):
        panel_label = _extract_csv_panel_label(filepath.stem, figure_number)
        with open(filepath, newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            if not reader.fieldnames:
                continue
            headers = list(reader.fieldnames)
            data = []
            for row in reader:
                converted = {}
                for k, v in row.items():
                    try:
                        converted[k] = float(v)
                    except (ValueError, TypeError):
                        converted[k] = v
                data.append(converted)

        if not data:
            continue

        has_errors = any("err" in h.lower() or "lower" in h.lower() or "upper" in h.lower()
                        for h in headers)
        plot_type = _infer_plot_type(headers, data, panel_label)

        panels.append(PanelSpec(
            sheet_name=panel_label,
            figure_number=figure_number,
            panel_label=_extract_panel_label(panel_label),
            columns=headers,
            data=data,
            plot_type=plot_type,
            has_error_bars=has_errors,
        ))
    return panels


def _extract_csv_panel_label(stem: str, figure_number: str) -> str:
    """Extract panel label from CSV filename like SourceData_Fig1__a or SourceData_ExtFig3."""
    m = re.search(r"__(.+)$", stem)
    if m:
        return m.group(1)
    return "main"


def _extract_panel_label(sheet_name: str) -> str:
    """Extract panel letter from sheet name like 'a', 'b_inset', 'c_upper'."""
    clean = sheet_name.strip().lower()
    m = re.match(r"^([a-z])(?:_|$)", clean)
    if m:
        return m.group(1)
    return clean


def _infer_plot_type(headers: list[str], data: list[dict], sheet_name: str) -> str:
    """Infer whether data should be a line plot, bar chart, box plot, etc."""
    x_col = headers[0]
    x_values = [d[x_col] for d in data if d.get(x_col) is not None]

    if "inset" in sheet_name.lower():
        if len(x_values) <= 5:
            return "bar_with_error"
        return "line_with_error"

    all_numeric_x = all(isinstance(v, (int, float)) for v in x_values)
    has_errors = any("err" in h.lower() or "lower" in h.lower() for h in headers)

    if any(kw in sheet_name.lower() for kw in ["median", "q1", "q3", "iqr"]):
        return "boxplot"

    lower_headers = [h.lower() for h in headers]
    if any("median" in h for h in lower_headers) and any("q1" in h or "q3" in h for h in lower_headers):
        return "boxplot"

    if all_numeric_x and len(x_values) > 10:
        if has_errors:
            return "line_with_error"
        return "line"

    if not all_numeric_x or len(x_values) <= 10:
        if has_errors:
            return "bar_with_error"
        num_series = len([h for h in headers[1:] if "err" not in h.lower()
                         and "lower" not in h.lower() and "upper" not in h.lower()])
        if num_series >= 2:
            return "grouped_bar"
        return "bar"

    if has_errors:
        return "line_with_error"
    return "line"


def _get_series_and_errors(panel: PanelSpec) -> list[dict]:
    """Extract series with their error columns from panel headers."""
    headers = panel.columns
    series = []
    i = 1
    while i < len(headers):
        h = headers[i]
        h_lower = h.lower()
        if "err" in h_lower or "lower" in h_lower or "upper" in h_lower:
            i += 1
            continue

        entry: dict = {"name": h, "col": h, "lower_err": None, "upper_err": None}

        for j in range(i + 1, min(i + 3, len(headers))):
            hj = headers[j].lower()
            if "lower" in hj or (hj.endswith("_err") and "lower" in hj):
                entry["lower_err"] = headers[j]
            elif "upper" in hj or (hj.endswith("_err") and "upper" in hj):
                entry["upper_err"] = headers[j]
            elif hj == f"{h.lower()}_lower_err":
                entry["lower_err"] = headers[j]
            elif hj == f"{h.lower()}_upper_err":
                entry["upper_err"] = headers[j]

        if entry["lower_err"] is None and entry["upper_err"] is None:
            for hh in headers:
                hh_lower = hh.lower()
                base = h.lower().replace("_mean", "").replace("_median", "")
                if hh_lower == f"{base}_lower_err":
                    entry["lower_err"] = hh
                elif hh_lower == f"{base}_upper_err":
                    entry["upper_err"] = hh

        series.append(entry)
        i += 1

    return series


NATURE_COLORS = {
    "green": "#2ca02c",
    "blue": "#1f77b4",
    "red": "#d62728",
    "purple": "#9467bd",
    "orange": "#ff7f0e",
    "NonAI": "#1f77b4",
    "AI": "#d62728",
    "nonai": "#1f77b4",
    "ai": "#d62728",
}

DEFAULT_COLORS = ["#1f77b4", "#d62728", "#2ca02c", "#9467bd", "#ff7f0e",
                  "#8c564b", "#e377c2", "#7f7f7f", "#bcbd22", "#17becf"]


def _get_color(name: str, idx: int) -> str:
    lower = name.lower().replace("_mean", "").replace("_median", "").replace("_y", "")
    if lower in NATURE_COLORS:
        return NATURE_COLORS[lower]
    return DEFAULT_COLORS[idx % len(DEFAULT_COLORS)]


def render_panel(panel: PanelSpec, output_path: Path, figsize: tuple = (6, 4)) -> Path:
    """Render a single panel to a PNG file."""
    fig, ax = plt.subplots(figsize=figsize)
    x_col = panel.columns[0]
    x_vals = [d[x_col] for d in panel.data]

    series_list = _get_series_and_errors(panel)

    if panel.plot_type == "boxplot":
        _render_boxplot(ax, panel, x_vals, series_list)
    elif panel.plot_type in ("bar", "grouped_bar", "bar_with_error"):
        _render_bar(ax, panel, x_vals, series_list)
    elif panel.plot_type in ("line", "line_with_error"):
        _render_line(ax, panel, x_vals, series_list)
    else:
        _render_line(ax, panel, x_vals, series_list)

    title = f"Figure {panel.figure_number}{panel.panel_label}"
    if panel.title:
        title += f": {panel.title}"
    ax.set_title(title, fontsize=11, fontweight="bold")

    if panel.x_label:
        ax.set_xlabel(panel.x_label)
    if panel.y_label:
        ax.set_ylabel(panel.y_label)

    handles, labels = ax.get_legend_handles_labels()
    if handles:
        ax.legend(frameon=False)

    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    plt.tight_layout()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(str(output_path), dpi=150, bbox_inches="tight")
    plt.close(fig)
    logger.info("Rendered %s → %s", title, output_path)
    return output_path


def _render_line(ax, panel: PanelSpec, x_vals: list, series_list: list[dict]):
    for idx, s in enumerate(series_list):
        color = _get_color(s["name"], idx)
        y = [d.get(s["col"]) for d in panel.data]

        numeric_x = all(isinstance(v, (int, float)) for v in x_vals)
        xp = x_vals if numeric_x else list(range(len(x_vals)))

        if s["lower_err"] and s["upper_err"]:
            lower = [d.get(s["lower_err"]) for d in panel.data]
            upper = [d.get(s["upper_err"]) for d in panel.data]
            valid = [(xi, yi, lo, hi) for xi, yi, lo, hi in zip(xp, y, lower, upper)
                     if yi is not None and lo is not None and hi is not None]
            if valid:
                xs, ys, los, his = zip(*valid)
                ax.plot(xs, ys, color=color, label=s["name"], linewidth=1.5)
                ax.fill_between(xs, los, his, alpha=0.2, color=color)
        else:
            valid = [(xi, yi) for xi, yi in zip(xp, y) if yi is not None]
            if valid:
                xs, ys = zip(*valid)
                ax.plot(xs, ys, color=color, label=s["name"], linewidth=1.5)

        if not numeric_x:
            ax.set_xticks(xp)
            ax.set_xticklabels(x_vals, rotation=45, ha="right")


def _render_bar(ax, panel: PanelSpec, x_vals: list, series_list: list[dict]):
    n_series = len(series_list)
    width = 0.7 / max(n_series, 1)
    x_pos = np.arange(len(x_vals))

    for idx, s in enumerate(series_list):
        color = _get_color(s["name"], idx)
        y = [d.get(s["col"], 0) or 0 for d in panel.data]
        offset = (idx - n_series / 2 + 0.5) * width

        if s["lower_err"] and s["upper_err"]:
            lower = [d.get(s["lower_err"], 0) or 0 for d in panel.data]
            upper = [d.get(s["upper_err"], 0) or 0 for d in panel.data]
            yerr_low = [max(0, yi - lo) for yi, lo in zip(y, lower)]
            yerr_high = [max(0, hi - yi) for yi, hi in zip(y, upper)]
            ax.bar(x_pos + offset, y, width, yerr=[yerr_low, yerr_high],
                   color=color, label=s["name"], capsize=3, edgecolor="white")
        else:
            ax.bar(x_pos + offset, y, width, color=color, label=s["name"], edgecolor="white")

    ax.set_xticks(x_pos)
    labels = [str(v) for v in x_vals]
    ax.set_xticklabels(labels, rotation=45 if max(len(l) for l in labels) > 5 else 0, ha="right")


def _render_boxplot(ax, panel: PanelSpec, x_vals: list, series_list: list[dict]):
    headers = panel.columns[1:]
    groups: dict[str, dict[str, str]] = {}

    for h in headers:
        h_lower = h.lower()
        for prefix in ("nonai", "ai", "non_ai"):
            if h_lower.startswith(prefix):
                base = prefix.replace("_", "")
                remainder = h_lower[len(prefix):].lstrip("_")
                if "median" in remainder:
                    groups.setdefault(base, {})["median"] = h
                elif "q1" in remainder and "1.5" not in remainder:
                    groups.setdefault(base, {})["q1"] = h
                elif "q3" in remainder and "1.5" not in remainder:
                    groups.setdefault(base, {})["q3"] = h
                elif "1.5iqr" in remainder or ("q1" in remainder and "1.5" in remainder):
                    groups.setdefault(base, {})["whisker_low"] = h
                elif "q3" in remainder and "1.5" in remainder:
                    groups.setdefault(base, {})["whisker_high"] = h
                break

    if not groups or all("q1" not in g for g in groups.values()):
        _render_bar(ax, panel, x_vals, series_list)
        return

    display_names = {"nonai": "NonAI", "ai": "AI"}
    n_groups = len(groups)
    positions_base = np.arange(len(x_vals))
    width = 0.3

    for g_idx, (group_name, cols) in enumerate(groups.items()):
        color = _get_color(group_name, g_idx)
        display = display_names.get(group_name, group_name)
        offset = (g_idx - n_groups / 2 + 0.5) * width
        positions = positions_base + offset

        for i, row in enumerate(panel.data):
            med = row.get(cols.get("median"))
            q1 = row.get(cols.get("q1"))
            q3 = row.get(cols.get("q3"))
            wl = row.get(cols.get("whisker_low"))
            wh = row.get(cols.get("whisker_high"))

            if med is None or q1 is None or q3 is None:
                continue

            bp = ax.bxp([{
                "med": med, "q1": q1, "q3": q3,
                "whislo": wl if wl is not None else q1,
                "whishi": wh if wh is not None else q3,
                "fliers": [],
            }], positions=[positions[i]], widths=width * 0.8,
                patch_artist=True, manage_ticks=False)

            for patch in bp["boxes"]:
                patch.set_facecolor(color)
                patch.set_alpha(0.7)
            for line in bp["medians"]:
                line.set_color("black")

        ax.plot([], [], color=color, label=display, linewidth=5, alpha=0.7)

    ax.set_xticks(positions_base)
    ax.set_xticklabels([str(v) for v in x_vals])


def replicate_figure(
    source_files: list[Path],
    figure_number: str,
    output_dir: Path,
    caption: str = "",
) -> FigureReplication:
    """Replicate all panels of a figure from its source data file(s)."""
    result = FigureReplication(
        figure_number=figure_number,
        source_file=str(source_files[0]),
    )

    try:
        panels = read_source_panels(source_files, figure_number)
    except Exception as e:
        result.status = "failed"
        result.message = f"Failed to read source data: {e}"
        return result

    if not panels:
        result.status = "no_data"
        result.message = "No plottable sheets found in source data"
        return result

    result.panels = panels
    fig_dir = output_dir / f"fig_{figure_number}"
    fig_dir.mkdir(parents=True, exist_ok=True)

    for panel in panels:
        safe_name = panel.sheet_name.replace("/", "_").replace(" ", "_")
        out_path = fig_dir / f"panel_{safe_name}.png"
        try:
            render_panel(panel, out_path)
            result.output_paths.append(str(out_path))
        except Exception as e:
            logger.warning("Failed to render %s panel %s: %s",
                          figure_number, panel.sheet_name, e)
            result.message += f"Panel {panel.sheet_name} failed: {e}; "

    result.status = "success" if result.output_paths else "failed"
    if result.status == "success":
        result.message = f"Rendered {len(result.output_paths)} panel(s)"

    return result


def replicate_all_figures(
    package_dir: str | Path,
    output_dir: str | Path,
    paper_figures: list[dict] | None = None,
) -> list[FigureReplication]:
    """Discover and replicate all figures with source data."""
    package_dir = Path(package_dir)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    source_map = discover_source_data(package_dir)
    logger.info("Found source data for %d figures: %s",
                len(source_map), list(source_map.keys()))

    results = []
    for fig_num, source_paths in sorted(source_map.items()):
        caption = ""
        if paper_figures:
            for pf in paper_figures:
                pf_num = str(pf.get("number", "")).strip()
                if pf_num.lstrip("E") == fig_num.lstrip("E") or pf_num == fig_num:
                    caption = pf.get("caption", "")
                    break

        result = replicate_figure(source_paths, fig_num, output_dir, caption)
        results.append(result)

    return results


def generate_summary_html(results: list[FigureReplication], output_dir: Path) -> Path:
    """Generate an HTML gallery of all replicated figure panels."""
    html_parts = [
        "<!DOCTYPE html><html><head>",
        "<meta charset='utf-8'>",
        "<title>Replicated Figures</title>",
        "<style>",
        "body { font-family: -apple-system, sans-serif; max-width: 1200px; margin: 0 auto; padding: 20px; }",
        "h1 { border-bottom: 2px solid #333; padding-bottom: 10px; }",
        ".figure-group { margin: 30px 0; padding: 20px; border: 1px solid #ddd; border-radius: 8px; }",
        ".figure-group h2 { margin-top: 0; color: #333; }",
        ".panels { display: flex; flex-wrap: wrap; gap: 15px; }",
        ".panel { text-align: center; }",
        ".panel img { max-width: 400px; border: 1px solid #eee; border-radius: 4px; }",
        ".panel .label { font-size: 12px; color: #666; margin-top: 5px; }",
        ".status-success { color: #2ca02c; }",
        ".status-failed { color: #d62728; }",
        ".summary { background: #f8f9fa; padding: 15px; border-radius: 8px; margin-bottom: 30px; }",
        "</style></head><body>",
        "<h1>Replicated Figures</h1>",
    ]

    success_count = sum(1 for r in results if r.status == "success")
    total_panels = sum(len(r.output_paths) for r in results)
    html_parts.append(
        f"<div class='summary'>"
        f"<strong>{success_count}/{len(results)}</strong> figures replicated, "
        f"<strong>{total_panels}</strong> total panels rendered</div>"
    )

    for result in results:
        status_class = f"status-{result.status}"
        html_parts.append(f"<div class='figure-group'>")
        html_parts.append(
            f"<h2>Figure {result.figure_number} "
            f"<span class='{status_class}'>({result.status})</span></h2>"
        )
        html_parts.append(f"<p>Source: <code>{Path(result.source_file).name}</code> — {result.message}</p>")

        if result.output_paths:
            html_parts.append("<div class='panels'>")
            for path_str in result.output_paths:
                p = Path(path_str)
                rel = p.relative_to(output_dir)
                label = p.stem.replace("panel_", "").replace("_", " ")
                html_parts.append(
                    f"<div class='panel'>"
                    f"<img src='{rel}' alt='Panel {label}'>"
                    f"<div class='label'>Panel {label}</div></div>"
                )
            html_parts.append("</div>")

        html_parts.append("</div>")

    html_parts.append("</body></html>")

    out = output_dir / "replicated_figures.html"
    out.write_text("\n".join(html_parts))
    return out
