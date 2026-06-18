from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

from replication_manager.paper import extract_paper_manifest


class DocumentFormatParsingTests(unittest.TestCase):
    def test_markdown_headings_extract_clean_title_tables_figures_and_claims(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            paper_path = Path(temp_dir) / "paper.md"
            paper_path.write_text(
                "\n".join(
                    [
                        "---",
                        "author: Example Researcher",
                        "---",
                        "# Clean Markdown Replication Study",
                        "",
                        "The replication rate was 88.5% across runs.",
                        "",
                        "## Table 1. Main Effects",
                        "| outcome | estimate |",
                        "| --- | ---: |",
                        "| turnout | 0.42 |",
                        "",
                        "### Figure 1. Coefficient Plot",
                    ]
                ),
                encoding="utf-8",
            )

            manifest = extract_paper_manifest(paper_path)

            self.assertEqual(manifest.title, "Clean Markdown Replication Study")
            self.assertEqual(len(manifest.tables), 1)
            self.assertEqual(manifest.tables[0].number, "1")
            self.assertEqual(manifest.tables[0].title, "Main Effects")
            self.assertEqual(len(manifest.figures), 1)
            self.assertEqual(manifest.figures[0].caption, "Coefficient Plot")
            raw_claims = [claim.raw for claim in manifest.numeric_claims]
            self.assertIn("0.42", raw_claims)
            self.assertIn("88.5%", raw_claims)

    def test_latex_extracts_title_captions_table_claims_and_escaped_percentages(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            paper_path = Path(temp_dir) / "paper.tex"
            paper_path.write_text(
                r"""
\documentclass{article}
\title{LaTeX Treatment Effects}
\author{Example Researcher}
\begin{document}
\maketitle

The replication rate was 94.4\% in the main specification.

\begin{table}
\caption{Main Estimates}
\begin{tabular}{lr}
Outcome & Estimate \\
Treatment & 0.32 \\
Control & -0.11 \\
\end{tabular}
\end{table}

\begin{figure}
\caption{Coefficient Plot}
\end{figure}
\end{document}
""",
                encoding="utf-8",
            )

            manifest = extract_paper_manifest(paper_path)

            self.assertEqual(manifest.title, "LaTeX Treatment Effects")
            self.assertEqual(len(manifest.tables), 1)
            self.assertEqual(manifest.tables[0].number, "1")
            self.assertEqual(manifest.tables[0].title, "Main Estimates")
            self.assertEqual(len(manifest.figures), 1)
            self.assertEqual(manifest.figures[0].caption, "Coefficient Plot")
            table_claims = [claim.raw for claim in manifest.tables[0].numeric_claims]
            self.assertIn("0.32", table_claims)
            self.assertIn("-0.11", table_claims)
            raw_claims = [claim.raw for claim in manifest.numeric_claims]
            self.assertIn("94.4%", raw_claims)

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx is not installed")
    def test_docx_extracts_paragraphs_and_tables_without_zip_binary_garbage(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as temp_dir:
            paper_path = Path(temp_dir) / "paper.docx"
            document = Document()
            document.add_paragraph("DOCX Replication Study")
            document.add_paragraph("The response rate was 66.7% in the survey sample.")
            document.add_paragraph("Table 1. Survey Effects")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Outcome"
            table.cell(0, 1).text = "Estimate"
            table.cell(1, 0).text = "Support"
            table.cell(1, 1).text = "1.25"
            document.add_paragraph("Figure 1. Survey Distribution")
            document.save(paper_path)

            manifest = extract_paper_manifest(paper_path)

            self.assertEqual(manifest.title, "DOCX Replication Study")
            self.assertEqual(len(manifest.tables), 1)
            self.assertEqual(manifest.tables[0].number, "1")
            self.assertEqual(manifest.tables[0].title, "Survey Effects")
            self.assertIn("Support", manifest.tables[0].body)
            self.assertEqual(len(manifest.figures), 1)
            self.assertEqual(manifest.figures[0].caption, "Survey Distribution")
            raw_claims = [claim.raw for claim in manifest.numeric_claims]
            self.assertIn("1.25", raw_claims)
            self.assertIn("66.7%", raw_claims)

            extracted_text = "\n".join(
                [
                    manifest.title,
                    *(table.body for table in manifest.tables),
                    *(claim.context for claim in manifest.numeric_claims),
                    *(figure.caption for figure in manifest.figures),
                ]
            )
            self.assertNotIn("PK", extracted_text)
            self.assertNotIn("[Content_Types]", extracted_text)


if __name__ == "__main__":
    unittest.main()
